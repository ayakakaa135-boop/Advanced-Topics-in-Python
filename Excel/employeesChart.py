import openpyxl
from pathlib import Path
from openpyxl.chart import Reference, BarChart, BarChart3D, LineChart

# ملف الإكسل
file_name = "employees_salaries.xlsx"
folder_path = r"C:\Users\haama\Desktop\PycharmProjects\Advanced Topics in Python\Excel"
file_path = Path(folder_path) / file_name

# فتح الملف
wb = openpyxl.load_workbook(file_path)
sheet = wb.active

# تحديد البيانات
data = Reference(sheet, min_col=2, max_col=2, min_row=1, max_row=6)
categories = Reference(sheet, min_col=1, max_col=1, min_row=2, max_row=6)

# ----- خيارات الرسم المختلفة -----

# 1- رسم أعمدة ثلاثي الأبعاد
chart = BarChart3D()
chart.add_data(data, titles_from_data=True)
chart.set_categories(categories)
chart.title = "3D Bar Chart of Salaries"
chart.x_axis.title = "Employee"
chart.y_axis.title = "Salary"

# 2- رسم خطي (لو تحبي الخط بدل الأعمدة)
# chart = LineChart()
# chart.add_data(data, titles_from_data=True)
# chart.set_categories(categories)
# chart.title = "Line Chart of Salaries"
# chart.x_axis.title = "Employee"
# chart.y_axis.title = "Salary"

# تغيير لون الأعمدة (مثال لكل الأعمدة نفس اللون)
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice

# لو أعمدة 3D ممكن التحكم باللون هكذا:
# for series in chart.series:
#     series.graphicalProperties = GraphicalProperties(solidFill="FF5733")  # لون برتقالي

# إضافة الرسم البياني للورقة
sheet.add_chart(chart, "D5")

# حفظ الملف
wb.save(file_path)
print("Chart updated in:", file_path)
